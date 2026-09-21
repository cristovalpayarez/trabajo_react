"""
/api/reportes — reportes de ventas en PDF y Excel.
"""

from datetime import datetime, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Venta, DetalleVenta, Usuario
from app.security import require_permission

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])


@router.get("/ventas/pdf")
def reporte_ventas_pdf(
    fecha: str = Query(default=None, description="Fecha YYYY-MM-DD"),
    fecha_inicio: str = Query(default=None),
    fecha_fin: str = Query(default=None),
    usuario_actual: dict = Depends(require_permission("ventas.reportes")),
    db: Session = Depends(get_db),
):
    from app.factura_pdf import generar_reporte_ventas_pdf

    if fecha:
        try:
            fecha_obj = datetime.strptime(fecha, "%Y-%m-%d").date()
            inicio = datetime.combine(fecha_obj, datetime.min.time())
            fin = inicio + timedelta(days=1)
            titulo = f"Reporte diario de ventas - {fecha_obj.strftime('%d/%m/%Y')}"
        except ValueError:
            raise HTTPException(status_code=400, detail={"mensaje": "Formato de fecha inválido"})
    elif fecha_inicio and fecha_fin:
        try:
            fi = datetime.strptime(fecha_inicio, "%Y-%m-%d")
            ff = datetime.strptime(fecha_fin, "%Y-%m-%d") + timedelta(days=1)
            inicio = fi
            fin = ff
            titulo = f"Reporte de ventas - {fi.strftime('%d/%m/%Y')} al {ff.strftime('%d/%m/%Y')}"
        except ValueError:
            raise HTTPException(status_code=400, detail={"mensaje": "Formato de fecha inválido"})
    else:
        hoy = datetime.now().date()
        inicio = datetime.combine(hoy, datetime.min.time())
        fin = inicio + timedelta(days=1)
        titulo = f"Reporte diario de ventas - {hoy.strftime('%d/%m/%Y')}"

    ventas = (
        db.query(Venta)
        .filter(Venta.fecha_venta >= inicio, Venta.fecha_venta < fin)
        .order_by(Venta.id.desc())
        .all()
    )

    items = []
    for v in ventas:
        productos = ", ".join(f"{d.item_nombre} x{d.cantidad}" for d in v.detalles)
        items.append({
            "id": v.id,
            "fecha": v.fecha_venta.strftime("%d/%m/%Y %H:%M") if v.fecha_venta else "-",
            "cliente": v.cliente_nombre or "N/A",
            "productos": productos or "Sin detalles",
            "cantidad_total": sum(d.cantidad for d in v.detalles),
            "total": float(v.total),
            "estado": v.estado,
        })

    total_general = sum(i["total"] for i in items)

    pdf_bytes = generar_reporte_ventas_pdf(titulo, items, total_general)

    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=reporte_ventas.pdf"},
    )


@router.get("/ventas/excel")
def reporte_ventas_excel(
    fecha: str = Query(default=None),
    fecha_inicio: str = Query(default=None),
    fecha_fin: str = Query(default=None),
    usuario_actual: dict = Depends(require_permission("ventas.reportes")),
    db: Session = Depends(get_db),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(status_code=500, detail={"mensaje": "openpyxl no está instalado"})

    from openpyxl.drawing.image import Image as XLImage

    from app.branding import (
        XL_AMBAR,
        XL_BLANCO,
        XL_BORDE,
        XL_CYAN,
        XL_FONDO,
        XL_PANEL,
        XL_PANEL_ALT,
        XL_ROJO,
        XL_SUAVE,
        XL_TEXTO,
        XL_VERDE,
        XL_VIOLET,
    )

    if fecha:
        try:
            fecha_obj = datetime.strptime(fecha, "%Y-%m-%d").date()
            inicio = datetime.combine(fecha_obj, datetime.min.time())
            fin = inicio + timedelta(days=1)
            titulo_hoja = f"Ventas {fecha_obj.strftime('%d-%m-%Y')}"
        except ValueError:
            raise HTTPException(status_code=400, detail={"mensaje": "Formato de fecha inválido"})
    elif fecha_inicio and fecha_fin:
        try:
            fi = datetime.strptime(fecha_inicio, "%Y-%m-%d")
            ff = datetime.strptime(fecha_fin, "%Y-%m-%d") + timedelta(days=1)
            inicio = fi
            fin = ff
            titulo_hoja = f"Ventas {fi.strftime('%d-%m-%Y')} al {ff.strftime('%d-%m-%Y')}"
        except ValueError:
            raise HTTPException(status_code=400, detail={"mensaje": "Formato de fecha inválido"})
    else:
        hoy = datetime.now().date()
        inicio = datetime.combine(hoy, datetime.min.time())
        fin = inicio + timedelta(days=1)
        titulo_hoja = f"Ventas {hoy.strftime('%d-%m-%Y')}"

    ventas = (
        db.query(Venta)
        .filter(Venta.fecha_venta >= inicio, Venta.fecha_venta < fin)
        .order_by(Venta.id.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = titulo_hoja[:31]  # Excel limita el nombre de hoja a 31 caracteres
    ws.sheet_view.showGridLines = False

    # ---------- Estilos base (paleta de la pagina) ----------
    fondo = PatternFill("solid", start_color=XL_FONDO, end_color=XL_FONDO)
    panel = PatternFill("solid", start_color=XL_PANEL, end_color=XL_PANEL)
    panel_alt = PatternFill("solid", start_color=XL_PANEL_ALT, end_color=XL_PANEL_ALT)
    lado_borde = Side(style="thin", color=XL_BORDE)
    borde = Border(left=lado_borde, right=lado_borde, top=lado_borde, bottom=lado_borde)

    fuente_base = Font(name="Segoe UI", size=10, color=XL_TEXTO)
    fuente_sub = Font(name="Segoe UI", size=10, color=XL_SUAVE)
    fuente_head = Font(name="Segoe UI", size=10, bold=True, color=XL_BLANCO)
    fuente_total = Font(name="Segoe UI", size=11, bold=True, color=XL_BLANCO)

    centro = Alignment(horizontal="center", vertical="center")
    izquierda = Alignment(horizontal="left", vertical="center", wrap_text=True)
    derecha = Alignment(horizontal="right", vertical="center")

    # ---------- Fondo oscuro en toda la zona del reporte ----------
    MAX_FILA = max(60, len(ventas) + 30)
    for fila in ws.iter_rows(min_row=1, max_row=MAX_FILA, min_col=1, max_col=7):
        for cell in fila:
            cell.fill = fondo

    # ---------- Logo de la pagina (imagen real) ----------
    from app.branding import generar_logo_png

    logo_bytes = generar_logo_png()
    img = XLImage(BytesIO(logo_bytes))
    escala = 0.55  # el PNG sale grande; se reduce a tamano de encabezado
    img.width = int(img.width * escala)
    img.height = int(img.height * escala)
    ws.add_image(img, "B2")

    # Titulo del reporte y fecha de generacion
    ws["E2"] = "REPORTE DE VENTAS"
    ws["E2"].font = Font(name="Segoe UI", size=16, bold=True, color=XL_BLANCO)
    ws["E2"].alignment = Alignment(horizontal="left", vertical="center")
    ws["E3"] = titulo_hoja.replace("Ventas ", "Periodo: ")
    ws["E3"].font = Font(name="Segoe UI", size=11, bold=True, color=XL_CYAN)
    ws["E3"].alignment = Alignment(horizontal="left", vertical="center")
    ws["E4"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["E4"].font = fuente_sub
    ws["E4"].alignment = Alignment(horizontal="left", vertical="center")

    # Linea neon (violeta -> cian) bajo el encabezado
    for col in range(2, 8):
        c = ws.cell(row=6, column=col)
        c.fill = PatternFill(
            "solid",
            start_color=XL_VIOLET if col <= 4 else XL_CYAN,
            end_color=XL_VIOLET if col <= 4 else XL_CYAN,
        )

    # ---------- Resumen en una linea (mas entendible que tarjetas) ----------
    total_general = sum(float(v.total) for v in ventas)
    promedio = total_general / len(ventas) if ventas else 0

    resumen = [
        ("Ventas del periodo:", str(len(ventas)), XL_BLANCO),
        ("Ingresos totales:", f"${total_general:,.2f}", XL_CYAN),
        ("Promedio por venta:", f"${promedio:,.2f}", XL_VIOLET),
    ]
    col0 = 2
    for etiqueta, valor, color in resumen:
        c_et = ws.cell(row=8, column=col0, value=etiqueta)
        c_et.font = fuente_sub
        c_et.alignment = Alignment(horizontal="left", vertical="center")
        c_va = ws.cell(row=8, column=col0 + 1, value=valor)
        c_va.font = Font(name="Segoe UI", size=11, bold=True, color=color)
        c_va.alignment = Alignment(horizontal="left", vertical="center")
        col0 += 2

    # ---------- Tabla de ventas ----------
    fila_tabla = 11
    headers = ["#", "Fecha", "Cliente", "Productos", "Cant.", "Total", "Estado"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=fila_tabla, column=col, value=header)
        cell.font = fuente_head
        cell.fill = panel
        cell.alignment = centro
        cell.border = Border(
            bottom=Side(style="thin", color=XL_CYAN),
            top=lado_borde, left=lado_borde, right=lado_borde,
        )

    estado_colores = {"Pagada": XL_VERDE, "Completada": XL_VERDE, "Pendiente": XL_AMBAR, "Emitida": XL_AMBAR, "Anulada": XL_ROJO, "Cancelada": XL_ROJO}

    for idx, v in enumerate(ventas, 1):
        productos = ", ".join(f"{d.item_nombre} x{d.cantidad}" for d in v.detalles)
        row_data = [
            v.id,
            v.fecha_venta.strftime("%d/%m/%Y %H:%M") if v.fecha_venta else "-",
            v.cliente_nombre or "N/A",
            productos or "Sin detalles",
            sum(d.cantidad for d in v.detalles),
            float(v.total),
            v.estado,
        ]
        fila = fila_tabla + idx
        relleno = panel_alt if idx % 2 else panel
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=fila, column=col, value=value)
            cell.font = fuente_base
            cell.fill = relleno
            cell.border = borde
            if col in (1, 5):
                cell.alignment = centro
            elif col == 6:
                cell.alignment = derecha
                cell.number_format = '"$"#,##0.00'
            else:
                cell.alignment = izquierda
        # Estado con color segun valor (verde/ambarto/rojo)
        c_estado = ws.cell(row=fila, column=7)
        color_estado = estado_colores.get(str(v.estado).capitalize(), XL_SUAVE)
        c_estado.font = Font(name="Segoe UI", size=10, bold=True, color=color_estado)
        c_estado.alignment = centro

    if not ventas:
        ws.merge_cells(start_row=fila_tabla + 1, start_column=1, end_row=fila_tabla + 1, end_column=7)
        c = ws.cell(row=fila_tabla + 1, column=1, value="No hay ventas registradas para el periodo seleccionado.")
        c.font = fuente_base
        c.alignment = centro

    # ---------- Fila de total general ----------
    fila_total = fila_tabla + max(len(ventas), 1) + 1
    ws.cell(row=fila_total, column=5, value="TOTAL:").font = fuente_total
    ws.cell(row=fila_total, column=5).alignment = derecha
    c2 = ws.cell(row=fila_total, column=6, value=total_general)
    c2.font = fuente_total
    c2.alignment = derecha
    c2.number_format = '"$"#,##0.00'
    for col in (5, 6):
        cc = ws.cell(row=fila_total, column=col)
        cc.fill = panel
        cc.border = Border(top=Side(style="thin", color=XL_VIOLET), bottom=lado_borde, left=lado_borde, right=lado_borde)

    # Pie con contacto
    fila_pie = fila_total + 2
    ws.merge_cells(start_row=fila_pie, start_column=1, end_row=fila_pie, end_column=7)
    c = ws.cell(
        row=fila_pie,
        column=1,
        value="NEXUS TECH · contacto@nexustech.com.co · WhatsApp +57 304 469 7238",
    )
    c.font = Font(name="Segoe UI", size=8, color=XL_SUAVE)
    c.alignment = centro

    # ---------- Anchos de columna y altura del encabezado ----------
    anchos = {"A": 8, "B": 20, "C": 24, "D": 44, "E": 9, "F": 14, "G": 13}
    for col_letra, ancho in anchos.items():
        ws.column_dimensions[col_letra].width = ancho
    for fila_num in (2, 3, 4, 5):
        ws.row_dimensions[fila_num].height = 20
    ws.row_dimensions[1].height = 8

    ws.freeze_panes = f"A{fila_tabla + 1}"


    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        iter([buffer.read()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=reporte_ventas.xlsx"},
    )
